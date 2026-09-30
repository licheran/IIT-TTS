""".xlsx containers for the workbook format (spec 03).

`export_xlsx` writes a workbook with frozen header rows and dropdowns on reference columns.
`import_xlsx` reads one in read-only mode and never evaluates formulas (cached values only), then
hands the raw sheets to the importer.
"""

import zipfile
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from tts.core.sheets import ColumnDef, Preset, SheetDef
from tts.io.importer import ImportIssue, ImportOutcome, RawRow, RawSheet, import_raw
from tts.io.tables import Table, WorkbookData, build_tables
from tts.presets import preset_for

MAX_BYTES = 20 * 1024 * 1024  # spec 06 section 9
_EXTRA_ROWS = 1000  # rows of headroom that keep their dropdowns


def export_xlsx(
    data: WorkbookData, dest: str | Path | BinaryIO, preset: Preset | None = None
) -> None:
    """Write a workbook: sheets in the preset's order, frozen headers, dropdowns."""
    preset = preset or preset_for(data.dataset)
    write_xlsx(build_tables(data, preset), dest, preset)


def write_xlsx(tables: Sequence[Table], dest: str | Path | BinaryIO, preset: Preset) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {s.name: s for s in preset.sheets}
    for table in tables:
        worksheet = workbook.create_sheet(table.name)
        for at, header in enumerate(table.headers, start=1):
            cell = worksheet.cell(row=1, column=at, value=header)
            cell.font = Font(bold=True)
        for r, row in enumerate(table.rows, start=2):
            for c, value in enumerate(row, start=1):
                if value is None:
                    continue
                cell = worksheet.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # never a formula, even if the text starts with "="
        worksheet.freeze_panes = "A2"
        for at, header in enumerate(table.headers, start=1):
            width = max(
                [len(header)]
                + [len(str(row[at - 1])) for row in table.rows if row[at - 1] is not None]
            )
            worksheet.column_dimensions[get_column_letter(at)].width = min(max(width + 2, 8), 48)
        sheet = sheets.get(table.name)
        if sheet is not None:
            _add_dropdowns(worksheet, table, sheet)
    workbook.save(dest)


def _add_dropdowns(worksheet: object, table: Table, sheet: SheetDef) -> None:
    last = len(table.rows) + 1 + _EXTRA_ROWS
    for at, header in enumerate(table.headers, start=1):
        column = sheet.column(header)
        formula = None if column is None else _dropdown(column)
        if formula is None:
            continue
        letter = get_column_letter(at)
        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="warning",  # a warning, so a value can still be typed on purpose
        )
        validation.add(f"{letter}2:{letter}{last}")
        worksheet.add_data_validation(validation)  # type: ignore[attr-defined]


def _dropdown(column: ColumnDef) -> str | None:
    """The list-validation formula for a column, or None when a dropdown does not fit."""
    if column.choices:
        return '"' + ",".join(column.choices) + '"'
    if column.kind == "str" and len(column.refs) == 1:
        target = column.refs[0]
        # The codes of the target sheet: column A from row 2 down to the last filled row.
        return f"OFFSET({target}!$A$2,0,0,MAX(1,COUNTA({target}!$A:$A)-1),1)"
    return None


def import_xlsx(
    source: str | Path | BinaryIO | bytes,
    preset: Preset | None = None,
    max_bytes: int = MAX_BYTES,
) -> ImportOutcome:
    """Read a workbook. Never raises for a bad file: the problem comes back as an issue."""
    stream = _open(source, max_bytes)
    if isinstance(stream, ImportIssue):
        return ImportOutcome((stream,))
    try:
        workbook = load_workbook(stream, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
        return ImportOutcome((ImportIssue("workbook", message="not a valid .xlsx file"),))
    try:
        raw = {
            ws.title: _raw_sheet(ws.title, ws.iter_rows(values_only=True))
            for ws in workbook.worksheets
        }
    finally:
        workbook.close()
    return import_raw(raw, preset)


def _open(source: str | Path | BinaryIO | bytes, max_bytes: int) -> BytesIO | ImportIssue:
    if isinstance(source, bytes):
        payload = source
    elif isinstance(source, str | Path):
        path = Path(source)
        if not path.is_file():
            return ImportIssue("workbook", message=f"file not found: {path}")
        if path.stat().st_size > max_bytes:
            return ImportIssue(
                "workbook", message=f"file is larger than {max_bytes // (1024 * 1024)} MB"
            )
        payload = path.read_bytes()
    else:
        payload = source.read(max_bytes + 1)
    if len(payload) > max_bytes:
        return ImportIssue(
            "workbook", message=f"file is larger than {max_bytes // (1024 * 1024)} MB"
        )
    return BytesIO(payload)


def _raw_sheet(name: str, rows: object) -> RawSheet:
    iterator = iter(rows)  # type: ignore[call-overload]
    first = next(iterator, ())
    sheet = RawSheet(name, ["" if v is None else str(v) for v in first])
    for number, values in enumerate(iterator, start=2):
        cells = list(values)
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in cells):
            continue
        sheet.rows.append(RawRow(number, cells))
    return sheet
