"""Template rows of a version 1 file (P9.2, ADR-0007): expanded once on import, never kept.

The application has no Templates table and no expand endpoint. A file that has template rows is
imported as activities, and the dataset then runs like any hand-made dataset.
"""

from collections.abc import Iterator
from pathlib import Path

import pytest
from api_helpers import create_dataset
from fastapi.testclient import TestClient

from tts.api.main import create_app
from tts.worker.runner import Worker

pytestmark = pytest.mark.integration

TEMPLATES = Path(__file__).resolve().parents[1] / "fixtures" / "l6" / "templates.xlsx"


@pytest.fixture
def client(db_url: str) -> Iterator[TestClient]:
    with TestClient(create_app(db_url)) as c:
        yield c


def upload(client: TestClient) -> int:
    dataset_id = create_dataset(client, "templates")
    files = {"file": ("templates.xlsx", TEMPLATES.read_bytes(), "application/octet-stream")}
    assert client.post(f"/datasets/{dataset_id}/import", files=files).json()["ok"]
    return dataset_id


def test_there_is_no_expand_endpoint(client: TestClient) -> None:
    dataset_id = create_dataset(client, "none")
    assert client.post(f"/datasets/{dataset_id}/expand").status_code == 404


def test_importing_template_rows_makes_the_activities_and_keeps_no_templates(
    client: TestClient,
) -> None:
    dataset_id = upload(client)
    assert client.get(f"/datasets/{dataset_id}").json()["kind"] == "hand_made"
    activities = client.get(f"/datasets/{dataset_id}/tables/Activities").json()
    assert activities["total"] == 77
    assert {r["values"]["template"] for r in activities["rows"]} == {None}
    names = [s["name"] for s in client.get(f"/datasets/{dataset_id}/schema").json()["sheets"]]
    assert "Templates" in names  # the file format still reads it ...
    assert client.get(f"/datasets/{dataset_id}/tables/Templates").status_code == 404  # ... no table


def test_the_export_of_an_imported_template_file_has_no_templates_sheet(client: TestClient) -> None:
    import io

    from openpyxl import load_workbook

    dataset_id = upload(client)
    exported = client.get(f"/datasets/{dataset_id}/export")
    names = load_workbook(io.BytesIO(exported.content)).sheetnames
    assert "Templates" not in names and "Activities" in names


def test_a_run_of_the_imported_activities_succeeds(client: TestClient, session_factory) -> None:
    dataset_id = upload(client)
    run_id = client.post(f"/datasets/{dataset_id}/runs", json={"time_limit_s": 30}).json()["run_id"]
    Worker(session_factory, "w").run_once()
    assert client.get(f"/runs/{run_id}").json()["status"] == "succeeded"
