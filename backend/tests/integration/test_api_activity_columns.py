"""A hand-made dataset's Activities table shows and edits its groups and teachers as columns.

ActivityGroups and ActivityTeachers stay in the file (version 1) but the application shows no table
for them (ADR-0007, Phase 23).
"""

import io
from collections.abc import Iterator

import pytest
from api_helpers import import_dataset, row_path
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tts.api.main import create_app

pytestmark = pytest.mark.integration


@pytest.fixture
def client(db_url: str) -> Iterator[TestClient]:
    with TestClient(create_app(db_url)) as c:
        yield c


@pytest.fixture
def dataset_id(client: TestClient, l6_dataset) -> int:  # type: ignore[no-untyped-def]
    return import_dataset(client, l6_dataset)


def activity(client: TestClient, dataset_id: int, code: str) -> dict:  # type: ignore[type-arg]
    rows = client.get(f"/datasets/{dataset_id}/tables/Activities").json()["rows"]
    return next(r["values"] for r in rows if r["key"] == code)


def test_the_activities_table_has_groups_and_teachers_columns(
    client: TestClient, dataset_id: int
) -> None:
    page = client.get(f"/datasets/{dataset_id}/tables/Activities").json()
    assert "groups" in page["headers"] and "teachers" in page["headers"]
    first = page["rows"][0]["values"]
    assert first["groups"] and first["teachers"]


def test_the_join_tables_are_not_tables_of_the_application(
    client: TestClient, dataset_id: int
) -> None:
    for name in ("ActivityGroups", "ActivityTeachers", "Templates"):
        assert client.get(f"/datasets/{dataset_id}/tables/{name}").status_code == 404


def test_changing_the_groups_cell_changes_who_attends_and_the_export_shows_it(
    client: TestClient, dataset_id: int
) -> None:
    rows = client.get(f"/datasets/{dataset_id}/tables/Activities").json()["rows"]
    code = rows[0]["key"]
    before = activity(client, dataset_id, code)
    kept = before["groups"].split(";")[0]
    changed = client.patch(
        row_path(dataset_id, "Activities", code), json={"values": {"groups": kept}}
    )
    assert changed.status_code == 200, changed.text
    assert activity(client, dataset_id, code)["groups"] == kept
    assert activity(client, dataset_id, code)["teachers"] == before["teachers"]

    exported = client.get(f"/datasets/{dataset_id}/export")
    sheet = load_workbook(io.BytesIO(exported.content))["ActivityGroups"]
    attending = [r[1].value for r in sheet.iter_rows(min_row=2) if r[0].value == code]
    assert attending == [kept]


def test_a_new_activity_can_be_typed_with_its_groups_and_teachers(
    client: TestClient, dataset_id: int
) -> None:
    template = client.get(f"/datasets/{dataset_id}/tables/Activities").json()["rows"][0]["values"]
    group = template["groups"].split(";")[0]
    teacher = template["teachers"].split(";")[0]
    made = client.post(
        f"/datasets/{dataset_id}/tables/Activities",
        json={
            "values": {
                "code": "NEW-01", "kind": "LEC", "duration": 2, "start_pattern": "2H",
                "delivery": "online", "groups": group, "teachers": teacher,
            }
        },
    )  # fmt: skip
    assert made.status_code == 201, made.text
    shown = activity(client, dataset_id, "NEW-01")
    assert (shown["groups"], shown["teachers"]) == (group, teacher)


def test_deleting_an_activity_takes_its_groups_and_teachers_with_it(
    client: TestClient, dataset_id: int
) -> None:
    code = client.get(f"/datasets/{dataset_id}/tables/Activities").json()["rows"][0]["key"]
    assert client.delete(row_path(dataset_id, "Activities", code)).status_code == 204
    exported = client.get(f"/datasets/{dataset_id}/export")
    book = load_workbook(io.BytesIO(exported.content))
    for name in ("ActivityGroups", "ActivityTeachers"):
        assert all(r[0].value != code for r in book[name].iter_rows(min_row=2))
