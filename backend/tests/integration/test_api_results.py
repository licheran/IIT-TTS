from collections.abc import Iterator
from io import BytesIO
from typing import Any

import pytest
from api_helpers import import_dataset
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tts.api.main import create_app
from tts.core.verifier import hard_violations, verify
from tts.io.workbook import import_xlsx
from tts.worker.runner import Worker

pytestmark = pytest.mark.integration

GROUP = "L6 SE / G1"


@pytest.fixture
def client(db_url: str) -> Iterator[TestClient]:
    with TestClient(create_app(db_url)) as c:
        yield c


@pytest.fixture
def solved(client: TestClient, session_factory, l6_dataset) -> dict[str, Any]:
    """Two finished runs of L6 (different seeds) on one dataset."""
    dataset_id = import_dataset(client, l6_dataset)
    worker = Worker(session_factory, "w")
    runs = []
    for seed in (1, 2):
        response = client.post(
            f"/datasets/{dataset_id}/runs",
            json={"seed": seed, "num_workers": 1, "time_limit_s": 30},
        )
        runs.append(response.json()["run_id"])
        assert worker.run_once()
    return {"dataset_id": dataset_id, "runs": runs}


def test_assignments_list_every_event_with_its_times_and_resources(
    client: TestClient, solved, l6_dataset
) -> None:
    rows = client.get(f"/runs/{solved['runs'][0]}/assignments").json()
    assert len(rows) == len(l6_dataset.events)
    one = rows[0]
    assert {"event", "kind", "day", "start_period", "end_period", "fixed", "chosen"} <= set(one)
    assert one["end_period"] != one["start_period"]  # every L6 event is two periods long


def test_the_grid_of_a_group_shows_each_event_as_one_cell(
    client: TestClient, solved, l6_expected: dict[str, Any]
) -> None:
    run_id = solved["runs"][0]
    grid = client.get(f"/runs/{run_id}/grid", params={"type": "StudentGroup", "code": GROUP}).json()
    assert grid["resource"] == GROUP
    assert sum(c["span"] for c in grid["cells"]) == l6_expected["group_periods_per_week"][GROUP]
    joint = [c for c in grid["cells"] if len(c["fixed"]) > 1]
    assert joint and all(GROUP in c["fixed"] for c in joint)


def test_a_grid_for_an_unknown_or_mistyped_resource_is_not_found(
    client: TestClient, solved
) -> None:
    run_id = solved["runs"][0]
    assert client.get(f"/runs/{run_id}/grid", params={"code": "nope"}).status_code == 404
    wrong = client.get(f"/runs/{run_id}/grid", params={"type": "Room", "code": GROUP})
    assert wrong.status_code == 404


def test_a_run_without_a_timetable_has_no_grid(client: TestClient, l6_dataset) -> None:
    dataset_id = import_dataset(client, l6_dataset)
    run_id = client.post(f"/datasets/{dataset_id}/runs", json={}).json()["run_id"]
    response = client.get(f"/runs/{run_id}/grid", params={"code": GROUP})
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "no_result"


def test_two_runs_differ_and_the_diff_names_the_moved_events(client: TestClient, solved) -> None:
    a, b = solved["runs"]
    diff = client.get(f"/runs/{a}/diff/{b}").json()
    assert diff["changes"], "different seeds should give different timetables"
    assert {c["kind"] for c in diff["changes"]} <= {"moved", "resources"}
    assert client.get(f"/runs/{a}/diff/{a}").json()["changes"] == []


def test_only_one_run_is_published_per_dataset(client: TestClient, solved) -> None:
    a, b = solved["runs"]
    assert client.post(f"/runs/{a}/publish").json()["published"] is True
    assert client.post(f"/runs/{b}/publish").json()["published"] is True
    published = [
        r["id"]
        for r in client.get(f"/datasets/{solved['dataset_id']}/runs").json()
        if r["published"]
    ]
    assert published == [b]


def test_a_run_that_did_not_finish_cannot_be_published(client: TestClient, l6_dataset) -> None:
    dataset_id = import_dataset(client, l6_dataset)
    run_id = client.post(f"/datasets/{dataset_id}/runs", json={}).json()["run_id"]
    response = client.post(f"/runs/{run_id}/publish")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "not_publishable"


def test_the_html_export_has_a_grid_for_the_group(client: TestClient, solved) -> None:
    run_id = solved["runs"][0]
    response = client.get(
        f"/runs/{run_id}/export", params={"format": "html", "type": "StudentGroup", "code": GROUP}
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    html = response.text
    assert "<h2>L6 SE / G1" in html and "Group</span>" in html
    assert html.count('class="event"') == len(
        client.get(f"/runs/{run_id}/grid", params={"code": GROUP}).json()["cells"]
    )


def test_the_html_export_of_a_type_has_one_grid_per_resource(client: TestClient, solved) -> None:
    run_id = solved["runs"][0]
    html = client.get(f"/runs/{run_id}/export", params={"type": "Teacher"}).text
    assert html.count("<h2>") > 10
    assert "Teacher</span>" in html


def test_the_xlsx_export_reads_back_with_its_assignments(
    client: TestClient, solved, l6_dataset
) -> None:
    run_id = solved["runs"][0]
    response = client.get(f"/runs/{run_id}/export", params={"format": "xlsx"})
    assert "Assignments" in load_workbook(BytesIO(response.content)).sheetnames
    outcome = import_xlsx(response.content)
    assert outcome.data is not None and outcome.data.result is not None
    assert outcome.data.dataset == l6_dataset
    assert hard_violations(verify(outcome.data.dataset, outcome.data.result)) == []


def test_the_csv_export_lists_the_events_of_a_resource(client: TestClient, solved) -> None:
    run_id = solved["runs"][0]
    everything = client.get(f"/runs/{run_id}/export", params={"format": "csv"}).text.splitlines()
    one = client.get(
        f"/runs/{run_id}/export", params={"format": "csv", "type": "StudentGroup", "code": GROUP}
    ).text.splitlines()
    assert everything[0].startswith("event,kind,reference,day")
    assert len(everything) == 78 and 1 < len(one) < len(everything)
