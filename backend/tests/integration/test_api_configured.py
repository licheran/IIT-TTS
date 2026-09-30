"""A configured dataset through the API: create, import, edit, pre-flight, run (ADR-0007)."""

import io

import pytest
from academic_config import dataset as configured_dataset
from api_helpers import create_dataset, row_path
from fastapi.testclient import TestClient

from tts.api.main import create_app
from tts.io.tables import WorkbookData
from tts.io.workbook import export_xlsx
from tts.store.repositories import RunRepo
from tts.worker.runner import Worker

pytestmark = pytest.mark.integration


@pytest.fixture
def client(db_url: str):
    with TestClient(create_app(db_url)) as c:
        yield c


def with_rooms(ds):  # type: ignore[no-untyped-def]
    """The configured dataset plus two labs, so a run has somewhere to put its sessions."""
    from academic_config import resource

    rooms = (
        resource("R1", "Room", capacity=100),
        resource("R2", "Room", capacity=100),
    )
    tagged = tuple(r.model_copy(update={"tags": (("room_type", "lab"),)}) for r in rooms)
    return ds.model_copy(update={"resources": (*ds.resources, *tagged)})


def workbook(ds) -> bytes:  # type: ignore[no-untyped-def]
    buffer = io.BytesIO()
    export_xlsx(WorkbookData(ds), buffer)
    return buffer.getvalue()


def upload(client: TestClient, dataset_id: int, payload: bytes):  # type: ignore[no-untyped-def]
    files = {"file": ("config.xlsx", payload, "application/octet-stream")}
    return client.post(f"/datasets/{dataset_id}/import", files=files)


def test_a_new_academic_dataset_is_configured_and_shows_the_version_two_sheets(client) -> None:
    dataset_id = create_dataset(client)
    assert client.get(f"/datasets/{dataset_id}").json()["kind"] == "configured"
    schema = client.get(f"/datasets/{dataset_id}/schema").json()
    names = [s["name"] for s in schema["sheets"]]
    assert (schema["kind"], schema["format_version"]) == ("configured", 2)
    assert "SessionTypes" in names
    for retired in ("Templates", "Activities", "ActivityGroups", "ActivityTeachers", "Pins"):
        assert retired not in names


def test_a_version_two_workbook_imports_and_its_rows_are_editable(client) -> None:
    dataset_id = create_dataset(client)
    response = upload(client, dataset_id, workbook(with_rooms(configured_dataset())))
    assert response.json()["ok"], response.json()
    page = client.get(f"/datasets/{dataset_id}/tables/SessionTypes").json()
    assert sorted(r["values"]["code"] for r in page["rows"]) == ["LEC", "TUT"]
    changed = client.patch(
        row_path(dataset_id, "SessionTypes", "LEC"), json={"values": {"max_groups": 2}}
    )
    assert changed.status_code == 200, changed.text
    again = client.get(f"/datasets/{dataset_id}/tables/SessionTypes").json()
    assert {r["values"]["code"]: r["values"]["max_groups"] for r in again["rows"]}["LEC"] == 2


def test_a_bad_version_two_workbook_reports_the_row_and_changes_nothing(client) -> None:
    from openpyxl import load_workbook

    dataset_id = create_dataset(client)
    sheets = load_workbook(io.BytesIO(workbook(with_rooms(configured_dataset()))))
    row = next(r for r in sheets["Teachers"].iter_rows(min_row=2) if r[0].value == "T3")
    headers = [c.value for c in sheets["Teachers"][1]]
    row[headers.index("modules")].value = "M3:TUT"
    out = io.BytesIO()
    sheets.save(out)
    response = upload(client, dataset_id, out.getvalue()).json()
    assert response["ok"] is False
    assert 'module "M3" has no session type "TUT"' in response["errors"][0]["message"]
    assert client.get(f"/datasets/{dataset_id}/tables/Teachers").json()["rows"] == []


def test_a_version_one_file_makes_the_dataset_hand_made_and_shows_the_version_one_sheets(
    client, l6_dataset
) -> None:
    from api_helpers import import_dataset

    dataset_id = import_dataset(client, l6_dataset)
    assert client.get(f"/datasets/{dataset_id}").json()["kind"] == "hand_made"
    schema = client.get(f"/datasets/{dataset_id}/schema").json()
    assert schema["format_version"] == 1
    assert "Activities" in [s["name"] for s in schema["sheets"]]


def test_the_export_of_a_configured_dataset_is_version_two_and_round_trips(client) -> None:
    dataset_id = create_dataset(client)
    assert upload(client, dataset_id, workbook(with_rooms(configured_dataset()))).json()["ok"]
    exported = client.get(f"/datasets/{dataset_id}/export")
    assert exported.status_code == 200
    second = create_dataset(client, "copy")
    assert upload(client, second, exported.content).json()["ok"]
    a = client.get(f"/datasets/{dataset_id}/tables/Modules").json()["rows"]
    b = client.get(f"/datasets/{second}/tables/Modules").json()["rows"]
    assert a == b


def test_the_table_editor_refuses_a_configuration_mistake(client) -> None:
    dataset_id = create_dataset(client)
    assert upload(client, dataset_id, workbook(with_rooms(configured_dataset()))).json()["ok"]
    changed = client.patch(
        row_path(dataset_id, "Teachers", "T1"), json={"values": {"modules": "M1;M9"}}
    )
    assert changed.status_code == 422
    assert 'unknown code "M9"' in changed.json()["error"]["details"][0]["message"]


def test_preflight_reports_a_configuration_mistake_found_in_stored_data(
    client, session_factory
) -> None:
    from tts.store.repositories import DatasetRepo

    dataset_id = create_dataset(client)
    assert upload(client, dataset_id, workbook(with_rooms(configured_dataset()))).json()["ok"]
    with session_factory() as session:
        repo = DatasetRepo(session)
        stored = repo.load(dataset_id)
        broken = stored.model_copy(
            update={
                "resources": tuple(
                    r.model_copy(update={"attributes": (("modules", "M1;M9"),)})
                    if r.code == "T1"
                    else r
                    for r in stored.resources
                )
            }
        )
        repo.save(dataset_id, broken)
        session.commit()
    issues = client.post(f"/datasets/{dataset_id}/preflight").json()["issues"]
    mistakes = [i for i in issues if i["kind"] == "configuration"]
    assert [i["severity"] for i in mistakes] == ["error"]
    assert mistakes[0]["message"] == 'Teachers "T1": unknown code "M9"'


def test_a_run_of_a_configured_dataset_creates_and_stores_its_sessions(
    client, session_factory
) -> None:
    dataset_id = create_dataset(client)
    assert upload(client, dataset_id, workbook(with_rooms(configured_dataset()))).json()["ok"]
    started = client.post(
        f"/datasets/{dataset_id}/runs", json={"num_workers": 1, "time_limit_s": 30}
    )
    assert started.status_code == 201, started.text
    run_id = started.json()["run_id"]
    assert Worker(session_factory, worker_id="w").run_once() is True
    run = client.get(f"/runs/{run_id}").json()
    assert run["status"] == "succeeded", run
    with session_factory() as session:
        result = RunRepo(session).result(run_id)
    codes = sorted(c.code for c in result.created)
    assert codes == ["M1-LEC-01", "M1-TUT-01", "M1-TUT-02", "M1-TUT-03", "M2-TUT-01", "M3-LEC-01"]
    rows = client.get(f"/runs/{run_id}/assignments").json()
    assert sorted(r["event"] for r in rows) == codes
    grid = client.get(f"/runs/{run_id}/grid", params={"code": "P1/G1"})
    assert grid.status_code == 200


def test_the_planned_sessions_list_each_module_and_kind_with_its_blocks(client) -> None:
    dataset_id = create_dataset(client)
    assert upload(client, dataset_id, workbook(with_rooms(configured_dataset()))).json()["ok"]
    planned = {p["demand"]: p for p in client.get(f"/datasets/{dataset_id}/sessions").json()}
    assert sorted(planned) == ["M1-LEC", "M1-TUT", "M2-TUT", "M3-LEC"]
    lecture, tutorial = planned["M1-LEC"], planned["M1-TUT"]
    assert lecture["groups"] == ["P1/G1", "P1/G2", "P2/G1"]
    assert (lecture["blocks"], lecture["sessions"], lecture["groups_per_session"]) == (1, 1, 3)
    assert (tutorial["blocks"], tutorial["per_week"], tutorial["sessions"]) == (3, 1, 3)
    assert planned["M2-TUT"]["groups"] == ["P1/G1"]


def test_a_hand_made_dataset_has_no_planned_sessions(client, l6_dataset) -> None:
    from api_helpers import import_dataset

    dataset_id = import_dataset(client, l6_dataset)
    assert client.get(f"/datasets/{dataset_id}/sessions").status_code == 409
