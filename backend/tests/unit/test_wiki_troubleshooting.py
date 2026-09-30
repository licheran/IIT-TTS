"""The troubleshooting pages (`docs/wiki/troubleshooting/`) cover every message the code gives."""

import re
from collections.abc import Callable
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tts.io.csvzip import import_csvzip
from tts.io.workbook import import_xlsx
from tts.store.repositories import TERMINAL

ROOT = Path(__file__).resolve().parents[3]
PAGES = ROOT / "docs" / "wiki" / "troubleshooting"
SRC = ROOT / "backend" / "src" / "tts"
L6 = ROOT / "backend" / "tests" / "fixtures" / "l6" / "l6.xlsx"
TEMPLATES = ROOT / "backend" / "tests" / "fixtures" / "l6" / "templates.xlsx"
CONFIG = ROOT / "backend" / "tests" / "fixtures" / "l6-config" / "l6-config.xlsx"


def page(name: str) -> str:
    return (PAGES / f"{name}.md").read_text("utf-8")


# --- Pre-flight issues -------------------------------------------------------------------------


def preflight_kinds() -> dict[str, str]:
    """Every kind `preflight/checks.py` can add, with its severity."""
    text = (SRC / "preflight" / "checks.py").read_text("utf-8")
    found = re.findall(r'self\.add\(\s*"(error|warning)",\s*"([a-z_]+)"', text)
    return {kind: severity for severity, kind in found}


def invariant_kinds() -> set[str]:
    text = (SRC / "core" / "model.py").read_text("utf-8")
    body = text[text.index("def validate_invariants") : text.index("def _matches_kind")]
    return set(re.findall(r'add\(\s*"([a-z_]+)",\s*"', body)) | {"duplicate_code"}


def test_the_check_finder_finds_the_known_kinds() -> None:
    assert {"no_candidate", "over_demand", "pooled_pressure", "empty_scope"} <= set(
        preflight_kinds()
    )
    assert {"hierarchy_cycle", "unknown_reference", "pooled_not_exclusive"} <= invariant_kinds()


def test_every_preflight_kind_has_an_entry_under_its_severity() -> None:
    text = page("preflight-issues")
    errors = text[text.index("## Errors") : text.index("## Warnings")]
    warnings = text[text.index("## Warnings") :].split("\n## Related")[0]
    for kind, severity in preflight_kinds().items():
        section = errors if severity == "error" else warnings
        assert f"`{kind}`" in section, f"{kind} ({severity})"


def test_every_data_model_issue_kind_is_listed() -> None:
    text = page("preflight-issues")
    for kind in invariant_kinds():
        assert f"| `{kind}` |" in text, kind


# --- Runs ---------------------------------------------------------------------------------------


def test_every_run_status_is_explained() -> None:
    text = page("run-statuses")
    for status in (*TERMINAL, "queued", "running"):
        assert f"`{status.replace('_', ' ')}`" in text, status


def test_every_run_diagnostic_kind_is_explained() -> None:
    text = page("run-statuses")
    source = (SRC / "worker" / "pipeline.py").read_text("utf-8")
    kinds = set(re.findall(r'Diagnostic\(\s*kind="([a-z_]+)"', source))
    assert {"no_solution", "unsupported_constraint", "infeasible_reason"} <= kinds
    explain = (SRC / "solver" / "explain.py").read_text("utf-8")
    kinds |= set(re.findall(r'Diagnostic\(\s*kind="([a-z_]+)"', explain))
    assert "infeasible_core" in kinds
    for kind in kinds:
        assert f"`{kind}`" in text, kind


def test_the_infeasible_page_covers_every_phrase_of_the_explanation() -> None:
    explain = (SRC / "solver" / "explain.py").read_text("utf-8")
    text = page("infeasible")
    for code_phrase, page_phrase in [
        ("has no allowed start for duration", "has no allowed start for duration"),
        ("matching", "matching"),
        ("pin of", "pin of"),
        ("no_overlap(", "no_overlap("),
        ("periods in all", "periods in all"),
        ("unavailable", "unavailable"),
        ("Conflicting rules: ", "Conflicting rules:"),
        ("the time grid and durations alone leave no solution", "the time grid and durations"),
        ("may include rules that are not needed", "may include rules that are not needed"),
    ]:
        assert code_phrase in explain, code_phrase
        assert page_phrase in text, page_phrase
    cli = (ROOT / "docs" / "cli.md").read_text("utf-8")
    assert (
        "Conflicting rules: Teacher THE unavailable Mon P01; pin of 6BUIS019C-LEC-01 to Mon P01"
        in cli
    )


def test_the_fallback_for_an_unexplained_infeasible_run_is_quoted() -> None:
    source = (SRC / "worker" / "pipeline.py").read_text("utf-8")
    assert "infeasible_unexplained" in source
    assert "could not be found in the time available" in page("infeasible")


# --- Import errors: each message is produced by the importer and quoted on the page -------------


def column_of(ws, name: str) -> int:
    return next(c.column for c in ws[1] if c.value == name)


def setc(sheet: str, row: int, column: str, value: object) -> Callable[[Workbook], None]:
    def change(wb: Workbook) -> None:
        ws = wb[sheet]
        ws.cell(row=row, column=column_of(ws, column)).value = value

    return change


def addrow(sheet: str, **values: object) -> Callable[[Workbook], None]:
    def change(wb: Workbook) -> None:
        ws = wb[sheet]
        row = ws.max_row + 1
        for key, value in values.items():
            ws.cell(row=row, column=column_of(ws, key)).value = value

    return change


def constraint(**values: object) -> Callable[[Workbook], None]:
    base = {"code": "X", "type": "max_days", "scope": "type:Teacher", "params": '{"max": 1}'}
    return addrow("Constraints", **{**base, **values})


def extra_column(sheet: str, name: str) -> Callable[[Workbook], None]:
    def change(wb: Workbook) -> None:
        ws = wb[sheet]
        ws.cell(row=1, column=ws.max_column + 1, value=name)

    return change


def drop_column(sheet: str, name: str) -> Callable[[Workbook], None]:
    def change(wb: Workbook) -> None:
        ws = wb[sheet]
        ws.delete_cols(column_of(ws, name))

    return change


def cycle(wb: Workbook) -> None:
    setc("Groups", 2, "parent", "L6 CS / G10")(wb)
    setc("Groups", 3, "parent", "L6 CS / G1")(wb)


def two_constraints(wb: Workbook) -> None:
    constraint()(wb)
    constraint(params='{"max": 2}')(wb)


CASES: list[tuple[str, Callable[[Workbook], None], Path, str]] = [
    ("no _meta", lambda wb: wb.remove(wb["_meta"]), L6, "_meta: sheet is missing"),
    (
        "format version",
        setc("_meta", 2, "value", "3"),
        L6,
        "_meta: format_version 3 not supported (max 2)",
    ),
    (
        "format version text",
        setc("_meta", 2, "value", "x"),
        L6,
        '_meta: format_version must be an integer, got "x"',
    ),
    ("no preset", lambda wb: wb["_meta"].delete_rows(3), L6, '_meta: "preset" is required'),
    ("unknown preset", setc("_meta", 3, "value", "nope"), L6, '_meta: unknown preset "nope"'),
    (
        "duplicate key",
        addrow("_meta", key="preset", value="academic_weekly"),
        L6,
        '_meta!R6C1 [key]: duplicate "preset"',
    ),
    ("unknown sheet", lambda wb: wb.create_sheet("Foo"), L6, "Foo: unknown sheet"),
    ("unknown column", extra_column("Rooms", "colour"), L6, "Rooms!R1C7 [colour]: unknown column"),
    (
        "missing column",
        drop_column("Rooms", "room_type"),
        L6,
        "Rooms!R1 [room_type]: missing column",
    ),
    (
        "duplicate column",
        extra_column("Rooms", "capacity"),
        L6,
        "Rooms!R1C7 [capacity]: duplicate column (first at C4)",
    ),
    ("blank required", setc("Rooms", 2, "room_type", None), L6, "Rooms!R2C5 [room_type]: required"),
    (
        "duplicate code",
        setc("Teachers", 3, "code", "AAM"),
        L6,
        'Teachers!R3C1 [code]: duplicate "AAM" (first at R2)',
    ),
    (
        "unknown code",
        setc("ActivityTeachers", 2, "teacher", "NOBODY"),
        L6,
        'ActivityTeachers!R2C2 [teacher]: unknown code "NOBODY"',
    ),
    (
        "bad integer",
        setc("Rooms", 2, "capacity", "thirty"),
        L6,
        'Rooms!R2C4 [capacity]: expected integer ≥ 0, got "thirty"',
    ),
    (
        "negative integer",
        setc("Groups", 2, "size", -1),
        L6,
        'Groups!R2C4 [size]: expected integer ≥ 0, got "-1"',
    ),
    (
        "zero duration",
        setc("Activities", 2, "duration", 0),
        L6,
        'Activities!R2C4 [duration]: expected integer ≥ 1, got "0"',
    ),
    (
        "bad time",
        setc("Periods", 2, "start", "8.30 am"),
        L6,
        'Periods!R2C2 [start]: expected HH:MM, got "8.30 am"',
    ),
    (
        "bad boolean",
        setc("Periods", 2, "is_break", "maybe"),
        L6,
        'Periods!R2C5 [is_break]: expected true or false, got "maybe"',
    ),
    (
        "end before start",
        setc("Periods", 2, "end", "08:00"),
        L6,
        "Periods!R2C3 [end]: end must be after start",
    ),
    (
        "bad choice",
        addrow("Availability", resource="AAM", day="Mon", period="P01", status="sometimes"),
        L6,
        'Availability!R2C4 [status]: expected one of unavailable, avoid, got "sometimes"',
    ),
    (
        "tag without equals",
        setc("Teachers", 2, "tags", "floor"),
        L6,
        'Teachers!R2C3 [tags]: expected key=value pairs separated by ";", got "floor"',
    ),
    (
        "duplicate tag",
        setc("Teachers", 2, "tags", "a=1;a=2"),
        L6,
        'Teachers!R2C3 [tags]: duplicate key "a"',
    ),
    (
        "bad json",
        constraint(params="{bad"),
        L6,
        "Constraints!R2C4 [params]: expected a JSON object, got invalid JSON "
        "(Expecting property name enclosed in double quotes)",
    ),
    (
        "json not an object",
        constraint(params="[1, 2]"),
        L6,
        "Constraints!R2C4 [params]: expected a JSON object",
    ),
    (
        "unknown type",
        constraint(type="max_gap"),
        L6,
        'Constraints!R2C2 [type]: "max_gap" is not in the catalogue',
    ),
    (
        "wrong scope kind",
        constraint(scope="kind:TUT"),
        L6,
        'Constraints!R2C3 [scope]: clause "kind:" does not select resources',
    ),
    (
        "bad selector",
        constraint(scope="teacher:"),
        L6,
        'Constraints!R2C3 [scope]: unknown clause "teacher:"',
    ),
    (
        "duplicate constraint",
        two_constraints,
        L6,
        'Constraints!R3C1 [code]: duplicate "X" (first at R2)',
    ),
    (
        "negative weight",
        constraint(weight=-1),
        L6,
        'Constraints!R2C6 [weight]: expected integer ≥ 0, got "-1"',
    ),
    (
        "hierarchy cycle",
        cycle,
        L6,
        "Groups!R2C3 [parent]: cycle L6 CS / G1 → L6 CS / G10 → L6 CS / G1",
    ),
    (
        "group without parent",
        setc("Groups", 2, "parent", None),
        L6,
        "Groups!R2C3 [parent]: required",
    ),
    (
        "unknown start pattern",
        setc("Activities", 2, "start_pattern", "9H"),
        L6,
        'Activities!R2C5 [start_pattern]: unknown code "9H"',
    ),
    (
        "unknown period in a list",
        setc("StartPatterns", 2, "start_periods", "P01;P99"),
        L6,
        'StartPatterns!R2C3 [start_periods]: unknown code "P99"',
    ),
    (
        "online with a room type",
        setc("Activities", 2, "delivery", "online"),
        L6,
        "Activities!R2C7 [room_type]: an online activity cannot have a room type",
    ),
    (
        "online with rooms",
        setc("Activities", 2, "delivery", "online"),
        L6,
        "Activities!R2C8 [room_count]: an online activity cannot have rooms",
    ),
    (
        "in person without a room type",
        setc("Activities", 2, "room_type", None),
        L6,
        "Activities!R2C7 [room_type]: required unless room_count is 0 (or the activity is online)",
    ),
    (
        "room count zero with a type",
        setc("Activities", 2, "room_count", 0),
        L6,
        "Activities!R2C8 [room_count]: expected integer ≥ 1 when room_type is set",
    ),
    (
        "batched without a size",
        setc("Templates", 2, "mode", "batched"),
        TEMPLATES,
        'Templates!R2C6 [batch_size]: required when mode is "batched"',
    ),
    (
        "bad template mode",
        setc("Templates", 2, "mode", "weekly"),
        TEMPLATES,
        'Templates!R2C4 [mode]: expected one of joint, per_group, batched, got "weekly"',
    ),
    (
        "bad template selector",
        setc("Templates", 2, "groups", "under:"),
        TEMPLATES,
        "Templates!R2C5 [groups]: under: needs a value",
    ),
    (
        "v2 unknown session type",
        setc("Modules", 2, "sessions", "LEC;XYZ"),
        CONFIG,
        'Modules!R2C6 [sessions]: unknown code "XYZ"',
    ),
    (
        "v2 setting without a value",
        setc("Modules", 2, "sessions", "LEC(max_groups)"),
        CONFIG,
        'Modules!R2C6 [sessions]: cannot read "LEC(max_groups)": expected name=value',
    ),
    (
        "v2 unknown setting",
        setc("Modules", 2, "sessions", "LEC(size=2)"),
        CONFIG,
        'Modules!R2C6 [sessions]: unknown setting "size" '
        "(known: start_pattern, delivery, room_type, max_groups, teachers, weekly)",
    ),
    (
        "v2 bad max_groups setting",
        setc("Modules", 2, "sessions", "LEC(max_groups=0)"),
        CONFIG,
        'Modules!R2C6 [sessions]: cannot read "LEC(max_groups=0)": '
        "max_groups must be an integer ≥ 1",
    ),
    (
        "v2 bad delivery setting",
        setc("Modules", 2, "sessions", "LEC(delivery=hybrid)"),
        CONFIG,
        'Modules!R2C6 [sessions]: cannot read "LEC(delivery=hybrid)": '
        "delivery must be in_person or online",
    ),
    (
        "v2 bad start pattern setting",
        setc("Modules", 2, "sessions", "LEC(start_pattern=9H)"),
        CONFIG,
        'Modules!R2C6 [sessions]: cannot read "LEC(start_pattern=9H)": unknown start pattern "9H"',
    ),
    (
        "v2 module without a level",
        setc("Modules", 2, "level", None),
        CONFIG,
        "Modules!R2C3 [level]: required",
    ),
    (
        "v2 teacher unknown module",
        setc("Teachers", 2, "modules", "NOPE"),
        CONFIG,
        'Teachers!R2C3 [modules]: unknown code "NOPE"',
    ),
    (
        "v2 teacher unknown kind",
        setc("Teachers", 2, "modules", "6BUIS019C:LAB"),
        CONFIG,
        'Teachers!R2C3 [modules]: "6BUIS019C:LAB": module "6BUIS019C" has no session type "LAB"',
    ),
    (
        "v2 option not optional",
        setc("Groups", 2, "options", "6COSC020C"),
        CONFIG,
        'Groups!R2C5 [options]: module "6COSC020C" is not optional',
    ),
    (
        "v2 option unknown",
        setc("Groups", 2, "options", "NOPE"),
        CONFIG,
        'Groups!R2C5 [options]: unknown code "NOPE"',
    ),
    (
        "v2 group under a group",
        setc("Groups", 2, "parent", "L6 CS / G10"),
        CONFIG,
        'Groups!R2C3 [parent]: unknown code "L6 CS / G10"',
    ),
    (
        "v2 online with a room type",
        setc("SessionTypes", 2, "delivery", "online"),
        CONFIG,
        "SessionTypes!R2C5 [room_type]: an online session cannot have a room type",
    ),
    (
        "v2 in person without a room type",
        setc("SessionTypes", 2, "room_type", None),
        CONFIG,
        'SessionTypes!R2C5 [room_type]: required when delivery is "in_person"',
    ),
    (
        "v2 no groups per session",
        setc("SessionTypes", 2, "max_groups", 0),
        CONFIG,
        'SessionTypes!R2C6 [max_groups]: expected integer ≥ 1, got "0"',
    ),
    (
        "v2 activities sheet",
        lambda wb: wb.create_sheet("Activities"),
        CONFIG,
        "Activities: unknown sheet",
    ),
]


@pytest.mark.parametrize(("label", "change", "source", "message"), CASES, ids=[c[0] for c in CASES])
def test_each_import_error_is_produced_and_quoted_on_the_page(
    label: str, change: Callable[[Workbook], None], source: Path, message: str
) -> None:
    wb = load_workbook(source)
    change(wb)
    buffer = BytesIO()
    wb.save(buffer)
    produced = [e.format() for e in import_xlsx(buffer.getvalue()).errors]
    assert message in produced, (label, produced[:5])
    assert f"`{message}`" in page("import-errors"), f"not on the page: {message}"


def test_the_file_level_messages_are_produced_and_quoted() -> None:
    text = page("import-errors")
    cases = [
        ([e.format() for e in import_xlsx(b"hello").errors], "workbook: not a valid .xlsx file"),
        ([e.format() for e in import_csvzip(b"hello").errors], "workbook: not a valid .zip file"),
        (
            [e.format() for e in import_xlsx(L6.read_bytes(), max_bytes=10).errors],
            "workbook: file is larger than 20 MB".replace("20", "0"),
        ),
    ]
    for produced, message in cases:
        assert message in produced, produced
    for quoted in ("workbook: not a valid .xlsx file", "workbook: not a valid .zip file"):
        assert f"`{quoted}`" in text
    assert "`workbook: file is larger than 20 MB`" in text


def test_params_that_are_wrong_pass_the_import_and_are_caught_by_preflight() -> None:
    wb = load_workbook(L6)
    constraint(type="max_gaps", params='{"mx": 1}')(wb)
    buffer = BytesIO()
    wb.save(buffer)
    assert import_xlsx(buffer.getvalue()).errors == ()
    assert "not caught at import" in page("import-errors")
