import io
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from tts.core.model import (
    CapacityRule,
    Dataset,
    Event,
    FixedRequirement,
    PooledRequirement,
    Resource,
    ResourceType,
    Result,
)

# L6 fixtures (l6_dataset, l6_locked_result) come from tests/fixtures/l6/conftest.py, which is
# only visible under that folder, so these tests build the dataset directly.
from tts.io.fet_html import parse_fet_groups_html, to_dataset
from tts.io.tables import Table, WorkbookData, WorkbookExportError, build_tables, row_key
from tts.io.workbook import export_xlsx, import_xlsx
from tts.presets import get_preset
from tts.presets.academic_weekly.preset import PRESET
from tts.presets.academic_weekly.sheets import SHEETS

L6_HTML = Path(__file__).resolve().parents[1] / "fixtures" / "l6" / "fet-groups-export.html"


@pytest.fixture(scope="module")
def l6() -> tuple[Dataset, Result]:
    return to_dataset(parse_fet_groups_html(L6_HTML))


@pytest.fixture(scope="module")
def l6_bytes(l6: tuple[Dataset, Result]) -> bytes:
    buffer = io.BytesIO()
    export_xlsx(WorkbookData(l6[0], l6[1], meta={"institution": "IIT"}, run="fet"), buffer)
    return buffer.getvalue()


def rows_of(workbook: Any, name: str) -> list[dict[str, Any]]:
    sheet = workbook[name]
    headers = [c.value for c in sheet[1]]
    return [
        dict(zip(headers, [c.value for c in row], strict=True))
        for row in sheet.iter_rows(min_row=2)
    ]


def test_sheets_come_in_the_specs_order_and_assignments_only_with_a_result(
    l6: tuple[Dataset, Result], l6_bytes: bytes
) -> None:
    full = load_workbook(io.BytesIO(l6_bytes)).sheetnames
    assert full == [s.name for s in SHEETS]
    buffer = io.BytesIO()
    export_xlsx(WorkbookData(l6[0]), buffer)
    assert load_workbook(io.BytesIO(buffer.getvalue())).sheetnames == [
        s.name for s in SHEETS if s.name != "Assignments"
    ]


def test_every_sheet_has_a_frozen_header_row(l6_bytes: bytes) -> None:
    workbook = load_workbook(io.BytesIO(l6_bytes))
    assert {ws.freeze_panes for ws in workbook.worksheets} == {"A2"}


def test_headers_are_the_defined_columns_without_convenience_columns(l6_bytes: bytes) -> None:
    workbook = load_workbook(io.BytesIO(l6_bytes))
    for sheet in SHEETS:
        expected = [c.name for c in sheet.columns if not c.convenience]
        assert [c.value for c in workbook[sheet.name][1]] == expected, sheet.name
    activities = [c.value for c in workbook["Activities"][1]]
    assert "groups" not in activities
    assert "teachers" not in activities


def test_meta_sheet(l6_bytes: bytes) -> None:
    rows = rows_of(load_workbook(io.BytesIO(l6_bytes)), "_meta")
    assert [(r["key"], r["value"]) for r in rows] == [
        ("format_version", "1"),
        ("preset", "academic_weekly"),
        ("institution", "IIT"),
    ]


def test_row_counts_and_typed_values(l6_bytes: bytes) -> None:
    workbook = load_workbook(io.BytesIO(l6_bytes))
    assert len(rows_of(workbook, "Groups")) == 30
    assert len(rows_of(workbook, "Teachers")) == 57
    assert len(rows_of(workbook, "Activities")) == 77
    periods = {r["code"]: r for r in rows_of(workbook, "Periods")}
    assert periods["P05"] == {
        "code": "P05",
        "start": "12:30",
        "end": "13:30",
        "order": 5,
        "is_break": True,
    }
    assert periods["P01"]["is_break"] is False
    assert isinstance(periods["P01"]["order"], int)
    pattern = rows_of(workbook, "StartPatterns")[0]
    assert (pattern["code"], pattern["duration"], pattern["start_periods"], pattern["days"]) == (
        "2H",
        2,
        "P01;P03;P06;P08;P10",
        None,
    )


def test_a_room_row(l6_bytes: bytes) -> None:
    rooms = {r["code"]: r for r in rows_of(load_workbook(io.BytesIO(l6_bytes)), "Rooms")}
    assert rooms["[2LA] -GP"] == {
        "code": "[2LA] -GP",
        "name": "[2LA] -GP",
        "building": "GP",
        "capacity": 90,
        "room_type": "lab",
        "tags": None,
    }
    assert rooms["Auditorium"]["capacity"] == 250
    assert rooms["Auditorium"]["room_type"] == "auditorium"


def test_activity_rows_carry_the_room_columns(l6_bytes: bytes) -> None:
    activities = {r["code"]: r for r in rows_of(load_workbook(io.BytesIO(l6_bytes)), "Activities")}
    online = next(r for r in activities.values() if r["delivery"] == "online")
    assert (online["room_type"], online["room_count"]) == (None, 0)
    physical = activities["6SENG010W-LEC-01"]
    assert (physical["delivery"], physical["room_type"], physical["room_count"]) == (
        "in_person",
        "auditorium",
        1,
    )
    assert (
        physical["module"],
        physical["kind"],
        physical["duration"],
        physical["start_pattern"],
    ) == (
        "6SENG010W",
        "LEC",
        2,
        "2H",
    )


def test_join_sheets_hold_the_fixed_resources(l6: tuple[Dataset, Result], l6_bytes: bytes) -> None:
    workbook = load_workbook(io.BytesIO(l6_bytes))
    groups = rows_of(workbook, "ActivityGroups")
    teachers = rows_of(workbook, "ActivityTeachers")
    kinds = {r.code: r.type for r in l6[0].resources}
    assert len(groups) == sum(kinds[f.resource] == "StudentGroup" for f in l6[0].fixed)
    assert len(teachers) == sum(kinds[f.resource] == "Teacher" for f in l6[0].fixed)
    assert {"activity": "6SENG010W-LEC-01", "group": "L6 SE / G1"} in groups


def test_assignment_rows_carry_derived_columns(l6_bytes: bytes) -> None:
    rows = {r["activity"]: r for r in rows_of(load_workbook(io.BytesIO(l6_bytes)), "Assignments")}
    row = rows["6SENG010W-LEC-01"]
    assert (row["run"], row["module"], row["kind"], row["day"]) == (
        "fet",
        "6SENG010W",
        "LEC",
        "Tue",
    )
    assert (row["start"], row["end"], row["rooms"], row["buildings"]) == (
        "08:30",
        "10:30",
        "Auditorium",
        "GP",
    )
    assert row["groups"] == "L6 SE / G1;L6 SE / G2;L6 SE / G3;L6 SE / G4;L6 SE / G5"
    assert len(row["teachers"].split(";")) == 5
    online = next(r for r in rows.values() if r["rooms"] is None)
    assert online["buildings"] is None


def test_dropdowns_are_on_reference_and_choice_columns(l6_bytes: bytes) -> None:
    workbook = load_workbook(io.BytesIO(l6_bytes))

    def dropdowns(sheet: str) -> dict[str, str]:
        ws = workbook[sheet]
        headers = {i: c.value for i, c in enumerate(ws[1], start=1)}
        found = {}
        for validation in ws.data_validations.dataValidation:
            first = str(validation.sqref).split(":")[0]
            index = column_index_from_string("".join(ch for ch in first if ch.isalpha()))
            found[headers[index]] = validation.formula1
            assert validation.type == "list"
            assert validation.errorStyle == "warning"
        return found

    activities = dropdowns("Activities")
    assert "StartPatterns!$A$2" in activities["start_pattern"]
    assert "Modules!$A$2" in activities["module"]
    assert "groups" not in activities
    assert dropdowns("ActivityTeachers")["teacher"].startswith("OFFSET(Teachers!$A$2")
    assert dropdowns("ActivityTeachers")["activity"].startswith("OFFSET(Activities!$A$2")
    assert dropdowns("Rooms")["building"].startswith("OFFSET(Buildings!$A$2")
    assert dropdowns("Assignments")["day"].startswith("OFFSET(Days!$A$2")
    assert "group" not in dropdowns("ActivityGroups")  # several target sheets: no dropdown
    assert "parent" not in dropdowns("Groups")


def test_choice_columns_get_a_literal_list(l6_bytes: bytes) -> None:
    ws = load_workbook(io.BytesIO(l6_bytes))["Templates"]
    formulas = [v.formula1 for v in ws.data_validations.dataValidation]
    assert '"joint,per_group,batched"' in formulas


def test_text_starting_with_an_equals_sign_is_never_a_formula() -> None:
    ds = Dataset(
        preset="academic_weekly",
        resource_types=PRESET.resource_types,
        resources=(Resource(code="=1+1", type="Teacher", name="+SUM(A1)"),),
    )
    buffer = io.BytesIO()
    export_xlsx(WorkbookData(ds), buffer)
    cell = load_workbook(io.BytesIO(buffer.getvalue()))["Teachers"]["A2"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"
    outcome = import_xlsx(buffer.getvalue())
    assert outcome.ok, [e.format() for e in outcome.errors]
    assert outcome.data.dataset.resources[0].code == "=1+1"  # type: ignore[union-attr]


def test_export_is_deterministic_apart_from_the_file_timestamp(l6: tuple[Dataset, Result]) -> None:
    data = WorkbookData(l6[0], l6[1])
    assert build_tables(data, PRESET) == build_tables(data, PRESET)


def test_notes_become_extra_columns_after_the_defined_ones(l6: tuple[Dataset, Result]) -> None:
    ds = l6[0]
    data = WorkbookData(
        ds,
        notes={
            ("Teachers", "HAWE"): {"x_comment": "head of dept"},
            ("Teachers", "HARR"): {"x_comment": "part time", "x_phone": "123"},
            ("Teachers", "GHOST"): {"x_ignored": "no such row"},
        },
    )
    table = next(t for t in build_tables(data, PRESET) if t.name == "Teachers")
    assert table.headers == ("code", "name", "tags", "x_comment", "x_phone")
    by_code = {row[0]: row for row in table.rows}
    assert by_code["HAWE"][3:] == ("head of dept", None)
    assert by_code["HARR"][3:] == ("part time", "123")
    assert by_code["AAM"][3:] == (None, None)


def test_a_workbook_without_notes_has_no_extra_columns(l6: tuple[Dataset, Result]) -> None:
    for table in build_tables(WorkbookData(l6[0]), PRESET):
        assert not any(h.startswith("x_") for h in table.headers)


def test_row_keys_identify_rows() -> None:
    teachers = next(s for s in SHEETS if s.name == "Teachers")
    assert row_key(teachers, {"code": "HAWE", "name": "x"}) == "HAWE"
    join = next(s for s in SHEETS if s.name == "ActivityTeachers")
    assert row_key(join, {"activity": "A1", "teacher": "T1"}) == "A1\x1fT1"
    availability = next(s for s in SHEETS if s.name == "Availability")
    assert (
        row_key(availability, {"resource": "T", "day": "Mon", "period": "*", "status": "avoid"})
        == "T\x1fMon\x1f*\x1favoid"
    )


# --- What the format cannot express --------------------------------------------------------------


def dataset_with(**changes: Any) -> Dataset:
    base = Dataset(
        preset="academic_weekly",
        resource_types=PRESET.resource_types,
        resources=(
            Resource(code="PR1", type="Programme"),
            Resource(code="G1", type="StudentGroup", parent="PR1", capacity=10),
            Resource(code="R1", type="Room", capacity=20, tags={"room_type": "lab"}),
        ),
        events=(Event(code="E1", kind="LEC", duration=1, start_pattern="1H"),),
        fixed=(FixedRequirement(event="E1", resource="G1"),),
    )
    return base.model_copy(update=changes)


def test_the_base_dataset_for_these_cases_exports() -> None:
    build_tables(WorkbookData(dataset_with()), PRESET)


def export_error(**changes: Any) -> str:
    with pytest.raises(WorkbookExportError) as caught:
        build_tables(WorkbookData(dataset_with(**changes)), PRESET)
    return str(caught.value)


def room_requirement(**kwargs: Any) -> PooledRequirement:
    defaults: dict[str, Any] = {
        "event": "E1",
        "resource_type": "Room",
        "filter": "tag:room_type=lab",
        "capacity_rule": CapacityRule.parse("sum_of_fixed:StudentGroup"),
    }
    return PooledRequirement(**{**defaults, **kwargs})


def test_a_standard_room_requirement_is_expressible() -> None:
    build_tables(WorkbookData(dataset_with(pooled=(room_requirement(),))), PRESET)


@pytest.mark.parametrize(
    ("pooled", "fragment"),
    [
        ((room_requirement(), room_requirement(ordinal=1)), "one pooled requirement"),
        ((room_requirement(ordinal=1),), "has no columns"),
        ((room_requirement(filter="tag:building=A"),), "does not test room_type"),
        ((room_requirement(filter="type:Room;tag:room_type=lab"),), "not a single tag test"),
        ((room_requirement(filter="tag:room_type!=lab"),), "not a single tag test"),
        ((room_requirement(filter="bogus"),), "not a single tag test"),
        ((room_requirement(capacity_rule=CapacityRule()),), "has no columns"),
        ((room_requirement(resource_type="Teacher"),), "has no columns"),
    ],
)
def test_a_pooled_requirement_the_columns_cannot_hold_is_refused(
    pooled: tuple[Any, ...], fragment: str
) -> None:
    assert fragment in export_error(pooled=pooled)


def test_a_resource_of_a_type_without_a_sheet_is_refused() -> None:
    types = (*PRESET.resource_types, ResourceType(code="Robot", exclusive=True))
    ds = dataset_with(
        resource_types=types,
        resources=(*dataset_with().resources, Resource(code="X1", type="Robot")),
    )
    with pytest.raises(WorkbookExportError, match='type "Robot" has no sheet'):
        build_tables(WorkbookData(ds), PRESET)


def test_a_fixed_requirement_no_sheet_can_list_is_refused() -> None:
    message = export_error(fixed=(FixedRequirement(event="E1", resource="R1"),))
    assert '"E1" is fixed on "R1"' in message


def test_a_parent_from_the_wrong_sheet_is_refused() -> None:
    resources = (
        Resource(code="G1", type="StudentGroup", parent="R1", capacity=10),
        Resource(code="R1", type="Room", capacity=20, tags={"room_type": "lab"}),
    )
    assert "expected one of Programmes, Groups" in export_error(resources=resources)


def test_a_separator_inside_a_code_that_is_listed_is_refused() -> None:
    ds = dataset_with(
        events=(Event(code="E1", kind="LEC", duration=1, start_pattern="1H", tags={"a": "x;y"}),)
    )
    with pytest.raises(WorkbookExportError, match="cannot be written"):
        build_tables(WorkbookData(ds), PRESET)


def test_an_attribute_without_a_column_is_refused() -> None:
    ds = dataset_with(
        resources=(
            Resource(code="G1", type="StudentGroup", capacity=10, attributes={"colour": "red"}),
        )
    )
    with pytest.raises(WorkbookExportError, match="have no column"):
        build_tables(WorkbookData(ds), PRESET)


def test_a_missing_required_value_is_refused() -> None:
    ds = dataset_with(
        resources=(
            Resource(code="PR1", type="Programme"),
            Resource(code="G1", type="StudentGroup", parent="PR1", capacity=10),
            Resource(code="R1", type="Room", capacity=20),
        )
    )
    with pytest.raises(WorkbookExportError, match="room_type is required"):
        build_tables(WorkbookData(ds), PRESET)


def test_meta_may_not_override_what_the_exporter_writes() -> None:
    with pytest.raises(WorkbookExportError, match="written by the exporter"):
        build_tables(WorkbookData(dataset_with(), meta={"preset": "other"}), PRESET)


def test_availability_covering_every_period_is_written_as_a_star() -> None:
    from datetime import time

    from tts.core.model import Availability, Day, Period, TimeModel

    tm = TimeModel(
        days=(Day(code="d1", order=1),),
        periods=tuple(
            Period(code=f"p{i}", start=time(7 + i), end=time(8 + i), order=i) for i in (1, 2)
        ),
    )
    full = tuple(
        Availability(resource="G1", day="d1", period=p, status="avoid") for p in ("p1", "p2")
    )
    part = full[:1]
    for availability, expected in ((full, ["*"]), (part, ["p1"])):
        tables = build_tables(
            WorkbookData(dataset_with(time=tm, availability=availability)), PRESET
        )
        table = next(t for t in tables if t.name == "Availability")
        assert [row[2] for row in table.rows] == expected


def test_the_preset_is_found_from_the_dataset() -> None:
    buffer = io.BytesIO()
    export_xlsx(WorkbookData(dataset_with()), buffer)
    assert load_workbook(io.BytesIO(buffer.getvalue())).sheetnames[0] == "_meta"
    assert get_preset("academic_weekly") is PRESET
    assert isinstance(build_tables(WorkbookData(dataset_with()), PRESET)[0], Table)
