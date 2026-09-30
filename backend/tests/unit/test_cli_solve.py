"""`tts solve` and `tts validate`, with every exit code."""

from datetime import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from tts import cli
from tts.cli import app
from tts.core.model import (
    CapacityRule,
    Dataset,
    Day,
    Event,
    FixedRequirement,
    Period,
    PooledRequirement,
    Resource,
    StartPattern,
    TimeModel,
    Violation,
)
from tts.core.verifier import verify
from tts.io.csvzip import import_csvzip
from tts.io.tables import WorkbookData
from tts.io.workbook import export_xlsx, import_xlsx
from tts.presets.academic_weekly.preset import PRESET
from tts.solver.registry import COMPILERS

runner = CliRunner()
L6_HTML = Path(__file__).resolve().parents[1] / "fixtures" / "l6" / "fet-groups-export.html"


def import_to(tmp_path: Path, name: str, *extra: str) -> Path:
    out = tmp_path / name
    result = runner.invoke(app, ["import-fet", str(L6_HTML), "--out", str(out), *extra])
    assert result.exit_code == 0, result.output
    return out


def solve_args(source: Path, out: Path, *extra: str) -> list[str]:
    return ["solve", str(source), "--out", str(out), "--time-limit", "30", "--workers", "1", *extra]


def edited(source: Path, target: Path, change: object) -> Path:
    workbook = load_workbook(source)
    change(workbook)  # type: ignore[operator]
    workbook.save(target)
    return target


def test_solve_writes_a_verified_result_workbook(tmp_path: Path) -> None:
    source = import_to(tmp_path, "l6.xlsx")
    out = tmp_path / "solved.xlsx"
    result = runner.invoke(app, solve_args(source, out))
    assert result.exit_code == 0, result.output
    assert "Solver: optimal" in result.output
    assert "Verifier: 0 hard violation(s), 0 soft, 0 warning(s)." in result.output
    solved = import_xlsx(out).data
    assert solved is not None and solved.result is not None
    assert len(solved.result.assignments) == 77
    assert solved.run == "seed-0"
    original = import_xlsx(source).data
    assert original is not None and solved.dataset == original.dataset
    assert verify(solved.dataset, solved.result) == []


def test_solve_keeps_the_workbooks_extras_and_ignores_old_assignments(tmp_path: Path) -> None:
    source = import_to(tmp_path, "l6.xlsx", "--with-assignments")
    out = tmp_path / "solved.zip"
    result = runner.invoke(app, solve_args(source, out, "--seed", "3"))
    assert result.exit_code == 0, result.output
    assert "Assignments are ignored" in result.output
    data = import_csvzip(out).data
    assert data is not None and data.run == "seed-3"
    assert "Group size: 30" in data.meta["assumptions"]
    original = import_xlsx(source).data
    assert original is not None and original.result != data.result


def one_period_dataset(*, rooms: int, groups_per_event: bool) -> Dataset:
    """Two events in a single period. With `groups_per_event` each has its own group."""
    resources = [Resource(code="PR1", type="Programme")]
    resources += [
        Resource(code=f"g{i}", type="StudentGroup", parent="PR1", capacity=10) for i in (1, 2)
    ]
    resources += [
        Resource(code=f"R{i}", type="Room", capacity=10, tags=(("room_type", "lab"),))
        for i in range(rooms)
    ]
    return Dataset(
        preset="academic_weekly",
        resource_types=PRESET.resource_types,
        reference_types=PRESET.reference_types,
        resources=tuple(resources),
        time=TimeModel(
            days=(Day(code="Mon", order=1),),
            periods=(Period(code="P1", start=time(8), end=time(9), order=1),),
            start_patterns=(StartPattern(code="1H", duration=1, start_periods=("P1",)),),
        ),
        events=(
            Event(code="a", kind="LEC", duration=1, start_pattern="1H"),
            Event(code="b", kind="LEC", duration=1, start_pattern="1H"),
        ),
        fixed=(
            FixedRequirement(event="a", resource="g1"),
            FixedRequirement(event="b", resource="g2" if groups_per_event else "g1"),
        ),
        pooled=tuple(
            PooledRequirement(
                event=e,
                resource_type="Room",
                filter="tag:room_type=lab",
                capacity_rule=CapacityRule.parse("sum_of_fixed:StudentGroup"),
            )
            for e in ("a", "b")
        )
        if rooms
        else (),
    )


def test_solve_reports_an_impossible_timetable_with_exit_code_2(tmp_path: Path) -> None:
    """Pre-flight can't see this one: two events, one room, one period (only a pressure warning)."""
    source = tmp_path / "impossible.xlsx"
    export_xlsx(WorkbookData(one_period_dataset(rooms=1, groups_per_event=True)), source)
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, solve_args(source, out))
    assert result.exit_code == 2, result.output
    assert "WARNING pooled_pressure: Room tag:room_type=lab: 2 periods needed, 1 available" in (
        result.output
    )
    assert "no timetable exists" in result.output
    assert (
        'Explanation: Conflicting rules: a needs 1 Room matching "tag:room_type=lab"; '
        'b needs 1 Room matching "tag:room_type=lab"; no_overlap(R0)'
    ) in result.output
    assert not out.exists()


def test_solve_is_blocked_by_a_preflight_error_with_exit_code_4(tmp_path: Path) -> None:
    """Two events on one group in one period: pre-flight sees the over-demand and stops."""
    source = tmp_path / "blocked.xlsx"
    export_xlsx(WorkbookData(one_period_dataset(rooms=0, groups_per_event=False)), source)
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, solve_args(source, out))
    assert result.exit_code == 4, result.output
    assert "ERROR over_demand: Group g1: needs 2 periods, 1 available" in result.output
    assert "Pre-flight: 1 error(s), 0 warning(s)." in result.output
    assert "Solver:" not in result.output  # it never got as far as solving
    assert not out.exists()


def test_solve_reports_a_teacher_with_no_free_periods_before_solving(tmp_path: Path) -> None:
    def block_a_teacher(workbook: object) -> None:
        availability = workbook["Availability"]  # type: ignore[index]
        periods = [c.value for c in workbook["Periods"]["A"][1:]]  # type: ignore[index]
        for day in [c.value for c in workbook["Days"]["A"][1:]]:  # type: ignore[index]
            for period in periods:
                availability.append(["HAWE", day, period, "unavailable"])

    broken = edited(import_to(tmp_path, "l6.xlsx"), tmp_path / "broken.xlsx", block_a_teacher)
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, solve_args(broken, out))
    assert result.exit_code == 4
    assert "ERROR over_demand: Teacher HAWE: needs 16 periods, 0 available" in result.output
    assert not out.exists()


def test_solve_says_why_an_event_cannot_be_placed(tmp_path: Path) -> None:
    """A pin on a slot where the teacher is unavailable: pre-flight passes, compiling finds it."""

    def pin_onto_a_blocked_slot(workbook: object) -> None:
        workbook["Availability"].append(["THE", "Mon", "P01", "unavailable"])  # type: ignore[index]
        workbook["Pins"].append(["6BUIS019C-LEC-01", "Mon", "P01", None, "user"])  # type: ignore[index]

    broken = edited(
        import_to(tmp_path, "l6.xlsx"), tmp_path / "pinned.xlsx", pin_onto_a_blocked_slot
    )
    result = runner.invoke(app, solve_args(broken, tmp_path / "out.xlsx"))
    assert result.exit_code == 2, result.output
    assert "Pre-flight: 0 error(s), 0 warning(s)." in result.output
    assert "has no start left after availability and pins" in result.output
    assert (
        "Explanation: Conflicting rules: Teacher THE unavailable Mon P01; "
        "pin of 6BUIS019C-LEC-01 to Mon P01"
    ) in result.output


def test_solve_prints_preflight_warnings_and_carries_on(tmp_path: Path) -> None:
    def add_empty_scope(workbook: object) -> None:
        row = ["C1", "max_gaps", "type:Nothing", '{"max": 2}', False, 1, True]
        workbook["Constraints"].append(row)  # type: ignore[index]

    source = edited(import_to(tmp_path, "l6.xlsx"), tmp_path / "warn.xlsx", add_empty_scope)
    result = runner.invoke(app, solve_args(source, tmp_path / "out.xlsx"))
    assert result.exit_code == 0, result.output
    assert "WARNING empty_scope: C1: scope matches nothing" in result.output
    assert "Pre-flight: 0 error(s), 1 warning(s)." in result.output


def test_solve_reports_a_clean_preflight(tmp_path: Path) -> None:
    result = runner.invoke(app, solve_args(import_to(tmp_path, "l6.xlsx"), tmp_path / "out.xlsx"))
    assert result.exit_code == 0, result.output
    assert "Pre-flight: 0 error(s), 0 warning(s)." in result.output


def test_solve_refuses_a_hard_constraint_it_cannot_compile(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delitem(COMPILERS, "max_gaps")

    def add_hard(workbook: object) -> None:
        row = ["C1", "max_gaps", "type:StudentGroup", '{"max": 2}', True, 1, True]
        workbook["Constraints"].append(row)  # type: ignore[index]

    source = edited(import_to(tmp_path, "l6.xlsx"), tmp_path / "hard.xlsx", add_hard)
    result = runner.invoke(app, solve_args(source, tmp_path / "out.xlsx"))
    assert result.exit_code == 1
    assert 'hard constraint "C1"' in result.output
    assert "not supported by the solver yet" in result.output


def test_solve_warns_about_a_soft_constraint_it_cannot_compile(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delitem(COMPILERS, "max_gaps")

    def add_soft(workbook: object) -> None:
        row = ["C1", "max_gaps", "type:StudentGroup", '{"max": 2}', False, 1, True]
        workbook["Constraints"].append(row)  # type: ignore[index]

    source = edited(import_to(tmp_path, "l6.xlsx"), tmp_path / "soft.xlsx", add_soft)
    result = runner.invoke(app, solve_args(source, tmp_path / "out.xlsx"))
    assert result.exit_code == 0, result.output
    assert 'Warning: constraint "C1" (max_gaps) is not supported' in result.output


def test_solve_reports_workbook_problems_and_writes_nothing(tmp_path: Path) -> None:
    def break_capacity(workbook: object) -> None:
        workbook["Rooms"].cell(row=2, column=4, value="thirty")  # type: ignore[index]

    bad = edited(import_to(tmp_path, "l6.xlsx"), tmp_path / "bad.xlsx", break_capacity)
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, solve_args(bad, out))
    assert result.exit_code == 1
    assert 'Rooms!R2C4 [capacity]: expected integer ≥ 0, got "thirty"' in result.output
    assert not out.exists()


def test_solve_never_writes_a_result_the_verifier_rejects(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake = Violation(code="no_overlap", constraint_code="H1", severity="hard", message="clash")
    monkeypatch.setattr(cli, "verify", lambda dataset, result: [fake])
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, solve_args(import_to(tmp_path, "l6.xlsx"), out))
    assert result.exit_code == 3
    assert "Violation: clash" in result.output
    assert not out.exists()


def test_solve_needs_a_known_result_extension(tmp_path: Path) -> None:
    result = runner.invoke(app, solve_args(import_to(tmp_path, "l6.xlsx"), tmp_path / "out.csv"))
    assert result.exit_code == 2
    assert "use .xlsx or .zip" in result.output


def test_validate_accepts_the_original_placements(tmp_path: Path) -> None:
    source = import_to(tmp_path, "l6-full.xlsx", "--with-assignments")
    result = runner.invoke(app, ["validate", str(source)])
    assert result.exit_code == 0, result.output
    assert "Verifier: 0 hard violation(s), 0 soft, 0 warning(s)." in result.output


def test_validate_reports_hard_violations_with_exit_code_3(tmp_path: Path) -> None:
    def clash(workbook: object) -> None:
        sheet = workbook["Assignments"]  # type: ignore[index]
        headers = [c.value for c in sheet[1]]
        day, start = headers.index("day") + 1, headers.index("start") + 1
        first_row = [sheet.cell(row=2, column=c).value for c in (day, start)]
        for row in range(3, sheet.max_row + 1):  # put every event on the first one's slot
            sheet.cell(row=row, column=day, value=first_row[0])
            sheet.cell(row=row, column=start, value=first_row[1])

    source = import_to(tmp_path, "l6-full.xlsx", "--with-assignments")
    moved = edited(source, tmp_path / "moved.xlsx", clash)
    result = runner.invoke(app, ["validate", str(moved)])
    assert result.exit_code == 3, result.output
    assert "HARD H1:" in result.output
    assert "hard violation(s)" in result.output


def test_validate_needs_assignments(tmp_path: Path) -> None:
    result = runner.invoke(app, ["validate", str(import_to(tmp_path, "l6.xlsx"))])
    assert result.exit_code == 1
    assert "no Assignments sheet" in result.output
