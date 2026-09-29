import json
from pathlib import Path

from typer.testing import CliRunner

from tts.cli import app
from tts.core.model import Dataset, Result
from tts.core.verifier import verify
from tts.io.fet_html import parse_fet_groups_html, to_dataset

runner = CliRunner()
L6_HTML = Path(__file__).resolve().parents[1] / "fixtures" / "l6" / "fet-groups-export.html"


def test_help_lists_every_command() -> None:
    result = runner.invoke(app, ["--help"])
    assert result.exit_code == 0
    for command in ("solve", "validate", "import-fet", "export"):
        assert command in result.output


def test_import_fet_writes_a_json_that_loads_back_into_the_same_dataset(tmp_path: Path) -> None:
    out = tmp_path / "l6.json"
    result = runner.invoke(app, ["import-fet", str(L6_HTML), "--out", str(out)])
    assert result.exit_code == 0, result.output
    assert "30 groups and 77 events" in result.output
    assert "Anomaly: 6CCGD007C LEC on Thursday at 17:30" in result.output

    payload = json.loads(out.read_text("utf-8"))
    assert payload["format"] == "tts-import-json"
    assert len(payload["anomalies"]) == 1
    assert any("Group size: 30" in line for line in payload["assumptions"])

    dataset = Dataset.model_validate(payload["dataset"])
    placements = Result.model_validate(payload["result"])
    expected_dataset, expected_result = to_dataset(parse_fet_groups_html(L6_HTML))
    assert dataset == expected_dataset
    assert placements == expected_result
    assert dataset.validate_invariants() == []
    assert verify(dataset, placements) == []


def test_import_fet_accepts_the_short_output_option(tmp_path: Path) -> None:
    out = tmp_path / "short.json"
    assert runner.invoke(app, ["import-fet", str(L6_HTML), "-o", str(out)]).exit_code == 0
    assert out.exists()


def test_import_fet_rejects_a_missing_input_file(tmp_path: Path) -> None:
    result = runner.invoke(
        app, ["import-fet", str(tmp_path / "nope.html"), "--out", str(tmp_path / "x.json")]
    )
    assert result.exit_code == 2


def test_import_fet_reports_unreadable_input_and_writes_nothing(tmp_path: Path) -> None:
    bad = tmp_path / "bad.html"
    bad.write_text("<html><body><p>not a timetable</p></body></html>", "utf-8")
    out = tmp_path / "out.json"
    result = runner.invoke(app, ["import-fet", str(bad), "--out", str(out)])
    assert result.exit_code == 1
    assert "no group tables found" in result.output
    assert not out.exists()


# --- Workbook output and export (P3.7) ---------------------------------------------------------


def import_to(tmp_path: Path, name: str, *extra: str) -> Path:
    out = tmp_path / name
    result = runner.invoke(app, ["import-fet", str(L6_HTML), "--out", str(out), *extra])
    assert result.exit_code == 0, result.output
    return out


def test_import_fet_writes_an_xlsx_workbook_that_reads_back(tmp_path: Path) -> None:
    from tts.io.workbook import import_xlsx

    out = import_to(tmp_path, "l6.xlsx")
    outcome = import_xlsx(out)
    assert outcome.ok, [e.format() for e in outcome.errors]
    data = outcome.data
    assert data is not None
    expected, _ = to_dataset(parse_fet_groups_html(L6_HTML))
    assert data.dataset == expected
    assert data.result is None  # the configuration only, unless asked
    assert data.meta["institution"] == "Informatics Institute of Technology"
    assert "Group size: 30" in data.meta["assumptions"]


def test_import_fet_can_include_the_original_placements(tmp_path: Path) -> None:
    from tts.io.workbook import import_xlsx

    data = import_xlsx(import_to(tmp_path, "l6-full.xlsx", "--with-assignments")).data
    assert data is not None and data.result is not None
    assert data.run == "fet"
    assert len(data.result.assignments) == 77
    assert verify(data.dataset, data.result) == []


def test_import_fet_writes_a_csv_zip(tmp_path: Path) -> None:
    from tts.io.csvzip import import_csvzip

    outcome = import_csvzip(import_to(tmp_path, "l6.zip"))
    assert outcome.ok
    assert outcome.data is not None and len(outcome.data.dataset.events) == 77


def test_import_fet_needs_a_known_output_extension(tmp_path: Path) -> None:
    result = runner.invoke(app, ["import-fet", str(L6_HTML), "--out", str(tmp_path / "out.txt")])
    assert result.exit_code == 2
    assert "use .xlsx or .zip" in result.output


def test_export_converts_between_formats_without_losing_anything(tmp_path: Path) -> None:
    from tts.io.csvzip import import_csvzip
    from tts.io.workbook import import_xlsx

    source = import_to(tmp_path, "l6.xlsx", "--with-assignments")
    as_zip = tmp_path / "copy.zip"
    result = runner.invoke(app, ["export", str(source), "--out", str(as_zip)])
    assert result.exit_code == 0, result.output
    back = tmp_path / "back.xlsx"
    assert runner.invoke(app, ["export", str(as_zip), "--out", str(back)]).exit_code == 0
    first, second, third = import_xlsx(source), import_csvzip(as_zip), import_xlsx(back)
    assert first.data == second.data == third.data
    assert first.data is not None and first.data.result is not None


def test_export_from_an_import_json(tmp_path: Path) -> None:
    from tts.io.workbook import import_xlsx

    json_file = import_to(tmp_path, "l6.json")
    plain = tmp_path / "plain.xlsx"
    assert runner.invoke(app, ["export", str(json_file), "--out", str(plain)]).exit_code == 0
    full = tmp_path / "full.xlsx"
    args = ["export", str(json_file), "--out", str(full), "--with-assignments"]
    assert runner.invoke(app, args).exit_code == 0
    assert import_xlsx(plain).data.result is None  # type: ignore[union-attr]
    assert import_xlsx(full).data.result is not None  # type: ignore[union-attr]
    assert "Group size: 30" in import_xlsx(plain).data.meta["assumptions"]  # type: ignore[union-attr]


def test_export_reports_every_problem_of_a_broken_workbook(tmp_path: Path) -> None:
    import io

    from openpyxl import load_workbook

    good = import_to(tmp_path, "l6.xlsx")
    workbook = load_workbook(good)
    workbook["Rooms"].cell(row=2, column=4, value="thirty")
    workbook["Groups"].cell(row=2, column=4, value="big")
    buffer = io.BytesIO()
    workbook.save(buffer)
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(buffer.getvalue())
    out = tmp_path / "out.xlsx"
    result = runner.invoke(app, ["export", str(broken), "--out", str(out)])
    assert result.exit_code == 1
    assert 'Rooms!R2C4 [capacity]: expected integer ≥ 0, got "thirty"' in result.output
    assert 'Groups!R2C4 [size]: expected integer ≥ 0, got "big"' in result.output
    assert "2 problem(s) found" in result.output
    assert not out.exists()


def test_export_rejects_unknown_extensions_and_bad_json(tmp_path: Path) -> None:
    good = import_to(tmp_path, "l6.xlsx")
    assert (
        runner.invoke(app, ["export", str(good), "--out", str(tmp_path / "x.csv")]).exit_code == 2
    )
    text = tmp_path / "input.txt"
    text.write_text("x", "utf-8")
    assert (
        runner.invoke(app, ["export", str(text), "--out", str(tmp_path / "x.xlsx")]).exit_code == 2
    )
    bad = tmp_path / "bad.json"
    bad.write_text("{}", "utf-8")
    result = runner.invoke(app, ["export", str(bad), "--out", str(tmp_path / "x.xlsx")])
    assert result.exit_code == 1
    assert "is not an import .json" in result.output
