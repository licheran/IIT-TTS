"""The file containers: what they accept, what they refuse, and that a bad file never raises."""

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from tts.io import csvzip
from tts.io.csvzip import export_csvzip, import_csvzip
from tts.io.fet_html import parse_fet_groups_html, to_dataset
from tts.io.tables import WorkbookData
from tts.io.workbook import export_xlsx, import_xlsx
from tts.presets.academic_weekly.sheets import SHEETS

L6_HTML = Path(__file__).resolve().parents[1] / "fixtures" / "l6" / "fet-groups-export.html"


@pytest.fixture(scope="module")
def data() -> WorkbookData:
    dataset, result = to_dataset(parse_fet_groups_html(L6_HTML))
    return WorkbookData(dataset, result)


def xlsx_bytes(data: WorkbookData) -> bytes:
    buffer = io.BytesIO()
    export_xlsx(data, buffer)
    return buffer.getvalue()


def zip_bytes(data: WorkbookData) -> bytes:
    buffer = io.BytesIO()
    export_csvzip(data, buffer)
    return buffer.getvalue()


def texts(outcome: object) -> list[str]:
    return [e.format() for e in outcome.errors]  # type: ignore[attr-defined]


# --- .xlsx ------------------------------------------------------------------------------------


def test_xlsx_can_be_read_from_a_path_a_string_a_stream_or_bytes(
    data: WorkbookData, tmp_path: Path
) -> None:
    path = tmp_path / "wb.xlsx"
    export_xlsx(data, path)
    content = path.read_bytes()
    for source in (path, str(path), io.BytesIO(content), content):
        outcome = import_xlsx(source)
        assert outcome.ok, texts(outcome)
        assert outcome.data == data


def test_a_file_that_is_not_a_workbook_is_reported_not_raised() -> None:
    assert texts(import_xlsx(b"this is not a spreadsheet")) == ["workbook: not a valid .xlsx file"]
    assert texts(import_xlsx(b"")) == ["workbook: not a valid .xlsx file"]
    other_zip = io.BytesIO()
    with zipfile.ZipFile(other_zip, "w") as archive:
        archive.writestr("hello.txt", "hi")
    assert texts(import_xlsx(other_zip.getvalue())) == ["workbook: not a valid .xlsx file"]


def test_a_missing_file_is_reported(tmp_path: Path) -> None:
    outcome = import_xlsx(tmp_path / "nope.xlsx")
    assert texts(outcome) == [f"workbook: file not found: {tmp_path / 'nope.xlsx'}"]
    assert outcome.data is None


def test_the_size_limit_applies_to_paths_bytes_and_streams(
    data: WorkbookData, tmp_path: Path
) -> None:
    content = xlsx_bytes(data)
    path = tmp_path / "wb.xlsx"
    path.write_bytes(content)
    for source in (path, content, io.BytesIO(content)):
        assert texts(import_xlsx(source, max_bytes=1000)) == ["workbook: file is larger than 0 MB"]


def test_formulas_are_never_evaluated() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    meta = workbook.create_sheet("_meta")
    meta.append(["key", "value"])
    meta.append(["format_version", 1])
    meta.append(["preset", "academic_weekly"])
    teachers = workbook.create_sheet("Teachers")
    teachers.append(["code", "name"])
    teachers["A2"] = "=1+1"  # a real formula, with no cached value
    teachers["B2"] = "Ann"
    buffer = io.BytesIO()
    workbook.save(buffer)
    outcome = import_xlsx(buffer.getvalue())
    assert texts(outcome) == ["Teachers!R2C1 [code]: required"]


def test_helper_sheets_are_ignored_and_blank_rows_skipped() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    meta = workbook.create_sheet("_meta")
    for row in (["key", "value"], ["format_version", 1], ["preset", "academic_weekly"]):
        meta.append(row)
    workbook.create_sheet("_lists").append(["anything"])
    teachers = workbook.create_sheet("Teachers")
    teachers.append(["code", "name"])
    teachers.append(["T1", "Ann"])
    teachers.append([None, None])
    teachers.append(["T2", "Bob"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    outcome = import_xlsx(buffer.getvalue())
    assert outcome.ok, texts(outcome)
    assert [r.code for r in outcome.data.dataset.resources] == ["T1", "T2"]  # type: ignore[union-attr]


def test_row_numbers_in_errors_count_skipped_blank_rows() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    meta = workbook.create_sheet("_meta")
    for row in (["key", "value"], ["format_version", 1], ["preset", "academic_weekly"]):
        meta.append(row)
    rooms = workbook.create_sheet("Rooms")
    rooms.append(["code", "capacity", "room_type"])
    rooms.append([None, None, None])
    rooms.append(["R1", "many", "lab"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert texts(import_xlsx(buffer.getvalue())) == [
        'Rooms!R3C2 [capacity]: expected integer ≥ 0, got "many"'
    ]


# --- CSV zip ----------------------------------------------------------------------------------


def test_the_zip_holds_one_utf8_csv_per_sheet_named_after_it(data: WorkbookData) -> None:
    with zipfile.ZipFile(io.BytesIO(zip_bytes(data))) as archive:
        assert archive.namelist() == [f"{s.name}.csv" for s in SHEETS if not s.import_only]
        teachers = archive.read("Teachers.csv").decode("utf-8")
        assert teachers.splitlines()[0] == "code,name,tags"
        assert not teachers.startswith("﻿")
        assert archive.read("_meta.csv").decode().splitlines()[:3] == [
            "key,value",
            "format_version,1",
            "preset,academic_weekly",
        ]


def test_the_zip_is_byte_for_byte_reproducible(data: WorkbookData) -> None:
    assert zip_bytes(data) == zip_bytes(data)


def test_csv_booleans_and_blanks(data: WorkbookData) -> None:
    with zipfile.ZipFile(io.BytesIO(zip_bytes(data))) as archive:
        periods = archive.read("Periods.csv").decode().splitlines()
        assert "P05,12:30,13:30,5,true" in periods
        assert "P01,08:30,09:30,1,false" in periods


def test_csv_can_be_read_from_a_path_a_string_a_stream_or_bytes(
    data: WorkbookData, tmp_path: Path
) -> None:
    path = tmp_path / "wb.zip"
    export_csvzip(data, path)
    content = path.read_bytes()
    for source in (path, str(path), io.BytesIO(content), content):
        outcome = import_csvzip(source)
        assert outcome.ok, texts(outcome)
        assert outcome.data == data


def test_a_bad_zip_is_reported_not_raised(tmp_path: Path) -> None:
    assert texts(import_csvzip(b"not a zip")) == ["workbook: not a valid .zip file"]
    assert texts(import_csvzip(tmp_path / "nope.zip"))[0].startswith("workbook: file not found")
    assert texts(
        import_csvzip(
            zip_bytes(WorkbookData(to_dataset(parse_fet_groups_html(L6_HTML))[0])), max_bytes=10
        )
    ) == ["workbook: file is larger than 0 MB"]


def make_zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buffer.getvalue()


META = b"key,value\nformat_version,1\npreset,academic_weekly\n"


def test_a_byte_order_mark_and_crlf_line_endings_are_accepted() -> None:
    files = {
        "_meta.csv": "﻿".encode() + META.replace(b"\n", b"\r\n"),
        "Teachers.csv": b"\xef\xbb\xbfcode,name\r\nT1,Ann\r\n",
    }
    outcome = import_csvzip(make_zip(files))
    assert outcome.ok, texts(outcome)
    assert [r.code for r in outcome.data.dataset.resources] == ["T1"]  # type: ignore[union-attr]


def test_other_zip_entries_are_ignored_and_folders_are_flattened() -> None:
    files = {
        "_meta.csv": META,
        "readme.txt": b"hello",
        "export/Teachers.csv": b"code,name\nT1,Ann\n",
    }
    outcome = import_csvzip(make_zip(files))
    assert outcome.ok, texts(outcome)
    assert outcome.summary == {"Teachers": 1}


def test_an_unknown_csv_and_non_utf8_text_are_reported() -> None:
    assert texts(import_csvzip(make_zip({"_meta.csv": META, "Scratch.csv": b"a\n1\n"}))) == [
        "Scratch: unknown sheet"
    ]
    assert texts(
        import_csvzip(make_zip({"_meta.csv": META, "Teachers.csv": b"code\n\xff\xfe\n"}))
    ) == ["Teachers: not valid UTF-8"]


def test_blank_lines_are_skipped_but_still_counted_in_row_numbers() -> None:
    files = {"_meta.csv": META, "Rooms.csv": b"code,capacity,room_type\n\nR1,many,lab\n"}
    assert texts(import_csvzip(make_zip(files))) == [
        'Rooms!R3C2 [capacity]: expected integer ≥ 0, got "many"'
    ]


def test_quoted_fields_with_commas_quotes_and_newlines_survive() -> None:
    files = {
        "_meta.csv": META,
        "Teachers.csv": b'code,name\nT1,"Smith, ""Jo""\nline two"\n',
    }
    outcome = import_csvzip(make_zip(files))
    assert outcome.data.dataset.resources[0].name == 'Smith, "Jo"\nline two'  # type: ignore[union-attr]


def test_a_zip_that_unpacks_to_too_much_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(csvzip, "MAX_UNPACKED_BYTES", 10)
    assert texts(import_csvzip(make_zip({"_meta.csv": META}))) == [
        "workbook: the zip unpacks to too much data"
    ]


def test_a_missing_meta_file_is_reported() -> None:
    assert texts(import_csvzip(make_zip({"Teachers.csv": b"code\nT1\n"}))) == [
        "_meta: sheet is missing"
    ]
