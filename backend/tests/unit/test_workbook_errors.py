"""P3.6: a workbook with five seeded errors of different kinds reports all five, and only those."""

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tts.io.fet_html import parse_fet_groups_html, to_dataset
from tts.io.tables import WorkbookData
from tts.io.workbook import export_xlsx, import_xlsx

L6_HTML = Path(__file__).resolve().parents[1] / "fixtures" / "l6" / "fet-groups-export.html"


@pytest.fixture(scope="module")
def clean_bytes() -> bytes:
    dataset, result = to_dataset(parse_fet_groups_html(L6_HTML))
    buffer = io.BytesIO()
    export_xlsx(WorkbookData(dataset, result), buffer)
    return buffer.getvalue()


def row_of(worksheet: object, code: str, column: int = 1) -> int:
    for row in worksheet.iter_rows(min_row=2):  # type: ignore[attr-defined]
        if row[column - 1].value == code:
            return int(row[column - 1].row)
    raise AssertionError(f"{code} not found")


def seeded(clean: bytes) -> tuple[bytes, dict[str, int]]:
    """The L6 workbook with five different mistakes. Returns it and the rows touched."""
    workbook = load_workbook(io.BytesIO(clean))
    rows: dict[str, int] = {}

    # 1. A number that is not a number.
    rooms = workbook["Rooms"]
    rows["room"] = row_of(rooms, "[2LA] -GP")
    rooms.cell(row=rows["room"], column=4, value="thirty")

    # 2. A teacher code with a typo.
    join = workbook["ActivityTeachers"]
    rows["join"] = 3
    join.cell(row=rows["join"], column=2, value="HAWEE")

    # 3. A duplicate activity code (a copy of the first data row, appended).
    activities = workbook["Activities"]
    rows["first_activity"] = 2
    duplicate = [c.value for c in activities[2]]
    rows["duplicate"] = activities.max_row + 1
    activities.append(duplicate)

    # 4. A constraint whose scope is not valid selector text.
    constraints = workbook["Constraints"]
    constraints.append(["C1", "max_gaps", "teacher:HAWE", '{"max": 2}', False, 5, True])
    rows["constraint"] = constraints.max_row

    # 5. Two groups that are each other's parent.
    groups = workbook["Groups"]
    first, second = row_of(groups, "L6 SE / G1"), row_of(groups, "L6 SE / G2")
    groups.cell(row=first, column=3, value="L6 SE / G2")
    groups.cell(row=second, column=3, value="L6 SE / G1")
    rows["cycle"] = first

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue(), rows


def test_all_five_seeded_errors_are_reported_with_sheet_row_column_and_message(
    clean_bytes: bytes,
) -> None:
    content, rows = seeded(clean_bytes)
    activities = load_workbook(io.BytesIO(content))["Activities"]
    code = activities.cell(row=rows["duplicate"], column=1).value
    first = rows["first_activity"]

    outcome = import_xlsx(content)

    assert not outcome.ok
    assert outcome.data is None
    assert [e.format() for e in outcome.errors] == [
        f"Groups!R{rows['cycle']}C3 [parent]: cycle L6 SE / G1 → L6 SE / G2 → L6 SE / G1",
        f'Rooms!R{rows["room"]}C4 [capacity]: expected integer ≥ 0, got "thirty"',
        f'Activities!R{rows["duplicate"]}C1 [code]: duplicate "{code}" (first at R{first})',
        f'ActivityTeachers!R{rows["join"]}C2 [teacher]: unknown code "HAWEE"',
        f'Constraints!R{rows["constraint"]}C3 [scope]: unknown clause "teacher:"',
    ]


def test_each_error_carries_its_position_as_data(clean_bytes: bytes) -> None:
    content, rows = seeded(clean_bytes)
    by_sheet = {e.sheet: e for e in import_xlsx(content).errors}
    assert (by_sheet["Rooms"].row, by_sheet["Rooms"].col, by_sheet["Rooms"].column) == (
        rows["room"],
        4,
        "capacity",
    )
    assert (by_sheet["ActivityTeachers"].row, by_sheet["ActivityTeachers"].col) == (rows["join"], 2)
    assert (by_sheet["Constraints"].row, by_sheet["Constraints"].col) == (rows["constraint"], 3)
    assert by_sheet["Activities"].message.startswith("duplicate ")
    assert by_sheet["Groups"].message.startswith("cycle ")


def test_fixing_the_errors_makes_the_same_workbook_import(clean_bytes: bytes) -> None:
    assert import_xlsx(clean_bytes).ok


def test_one_mistake_alone_reports_exactly_one_error(clean_bytes: bytes) -> None:
    workbook = load_workbook(io.BytesIO(clean_bytes))
    workbook["Rooms"].cell(row=2, column=4, value="thirty")
    out = io.BytesIO()
    workbook.save(out)
    (only,) = import_xlsx(out.getvalue()).errors
    assert only.sheet == "Rooms"
    assert only.row == 2
