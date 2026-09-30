"""L6 written as configuration (format version 2) end to end (P21.6, P22): the workbook imports,
passes pre-flight, solves, verifies and exports through the CLI and the API, and a session edited
through the API survives the rebuild."""

import io
from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from make_config import build  # type: ignore[import-not-found]
from typer.testing import CliRunner

from tts.api.expansion import prepare_with_preset
from tts.api.main import create_app
from tts.cli import app
from tts.core.verifier import hard_violations, verify
from tts.io.workbook import import_xlsx
from tts.store.db import make_engine, make_session_factory
from tts.store.repositories import DatasetRepo, RunRepo
from tts.worker.runner import Worker

HERE = Path(__file__).parent
WORKBOOK = HERE / "l6-config.xlsx"
runner = CliRunner()


def test_the_workbook_is_what_the_script_writes() -> None:
    outcome = import_xlsx(WORKBOOK)
    assert [e.format() for e in outcome.errors] == []
    assert outcome.data is not None and outcome.data.dataset == build()
    assert "assumptions" in outcome.data.meta


def test_the_workbook_holds_configuration_only() -> None:
    from openpyxl import load_workbook

    names = load_workbook(WORKBOOK).sheetnames
    assert "SessionTypes" in names
    for retired in ("Templates", "Activities", "ActivityGroups", "ActivityTeachers", "Pins"):
        assert retired not in names


def test_the_cli_solves_and_validates_the_configuration(tmp_path: Path) -> None:
    assert runner.invoke(app, ["preflight", str(WORKBOOK)]).exit_code == 0
    solved = tmp_path / "solved.xlsx"
    result = runner.invoke(
        app,
        ["solve", str(WORKBOOK), "--out", str(solved), "--time-limit", "120", "--workers", "4"],
    )
    assert result.exit_code == 0, result.output
    assert "Verifier: 0 hard violation(s)" in result.output


@pytest.fixture
def api(tmp_path: Path) -> Iterator[tuple[TestClient, object]]:
    url = f"sqlite:///{tmp_path / 'l6.db'}"
    with TestClient(create_app(url)) as client:
        yield client, url


def test_the_api_solves_the_configuration_and_keeps_an_edit_through_a_rebuild(
    api: tuple[TestClient, str],
) -> None:
    client, url = api
    created = client.post("/datasets", json={"name": "L6", "preset": "academic_weekly"}).json()
    dataset_id = created["id"]
    files = {"file": ("l6-config.xlsx", WORKBOOK.read_bytes(), "application/octet-stream")}
    assert client.post(f"/datasets/{dataset_id}/import", files=files).json()["ok"]
    assert client.post(f"/datasets/{dataset_id}/preflight").json()["has_errors"] is False
    planned = client.get(f"/datasets/{dataset_id}/sessions").json()
    sessions = sum(p["sessions"] for p in planned)
    assert sessions > 60

    factory = make_session_factory(make_engine(url))

    def solve() -> int:
        started = client.post(
            f"/datasets/{dataset_id}/runs", json={"num_workers": 2, "time_limit_s": 120}
        )
        assert started.status_code == 201, started.text
        assert Worker(factory, worker_id="w").run_once() is True
        run_id = int(started.json()["run_id"])
        assert client.get(f"/runs/{run_id}").json()["status"] == "succeeded"
        return run_id

    first = solve()
    table = client.get(f"/datasets/{dataset_id}/timetable").json()
    assert len(table["rows"]) == sessions and table["edits"] == 0
    target = next(r for r in table["rows"] if r["kind"] == "TUT")
    other_day = next(d for d in ("Mon", "Tue", "Wed", "Thu", "Fri") if d != target["day"])
    edited = client.put(
        f"/datasets/{dataset_id}/timetable/{target['code']}", json={"day": other_day}
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["edits"] == 1
    client.get(f"/datasets/{dataset_id}/timetable/check")  # judged at once; nothing stored

    second = solve()
    assert second != first
    kept = next(
        r
        for r in client.get(f"/datasets/{dataset_id}/timetable").json()["rows"]
        if r["code"] == target["code"]
    )
    assert kept["day"] == other_day and "day" in kept["edited"]
    with factory() as session:
        prepared = prepare_with_preset(DatasetRepo(session).load(dataset_id))
        result = RunRepo(session).result(second)
    assert hard_violations(verify(prepared, result)) == []

    exported = client.get(f"/datasets/{dataset_id}/export")
    assert exported.status_code == 200
    again = import_xlsx(io.BytesIO(exported.content))
    assert again.data is not None and again.data.dataset.kind == "configured"
